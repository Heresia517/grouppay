"""Service de generation des rapports mensuels."""
import calendar
from datetime import datetime
from decimal import Decimal
from io import BytesIO

from django.core.files.base import ContentFile
from django.db.models import Sum, Count, Avg, Q, F
from django.template.loader import render_to_string
from django.utils import timezone


def get_month_range(year, month):
    """Retourne (start_date, end_date) pour un mois."""
    start = datetime(year, month, 1)
    last_day = calendar.monthrange(year, month)[1]
    end = datetime(year, month, last_day, 23, 59, 59)
    return timezone.make_aware(start), timezone.make_aware(end)


def collect_monthly_data(organization, year, month):
    """Collecte toutes les donnees du mois."""
    from apps.orders.models import Order, OrderItem
    from apps.accounts.models import User

    start, end = get_month_range(year, month)

    # ─── Commandes ───
    orders_qs = Order.objects.filter(
        organization=organization,
        created_at__gte=start,
        created_at__lte=end,
    )
    orders_valid = orders_qs.filter(status__in=["VALIDEE", "EN_LIVRAISON", "LIVREE"])

    orders_data = {
        "total": orders_qs.count(),
        "validated": orders_valid.count(),
        "delivered": orders_qs.filter(status="LIVREE").count(),
        "rejected": orders_qs.filter(status="REFUSEE").count(),
        "pending": orders_qs.filter(status="EN_ATTENTE_VALIDATION").count(),
        "amount": orders_valid.aggregate(s=Sum("total_amount"))["s"] or Decimal("0"),
        "avg_amount": orders_valid.aggregate(a=Avg("total_amount"))["a"] or Decimal("0"),
    }

    # ─── Top produits ───
    top_products_qs = (
        OrderItem.objects.filter(
            order__in=orders_valid,
        )
        .values("product__name")
        .annotate(qty=Sum("quantity"), revenue=Sum(F("quantity") * F("unit_price")))
        .order_by("-qty")[:10]
    )
    top_products = [
        {"name": p["product__name"], "qty": p["qty"] or 0, "revenue": float(p["revenue"] or 0)}
        for p in top_products_qs
    ]

    # ─── Top membres ───
    top_members_qs = (
        orders_valid
        .values("member__email", "member__internal_id")
        .annotate(count=Count("id"), total=Sum("total_amount"))
        .order_by("-total")[:10]
    )
    top_members = [
        {
            "email": m["member__email"],
            "internal_id": m["member__internal_id"] or "—",
            "count": m["count"],
            "total": float(m["total"] or 0),
        }
        for m in top_members_qs
    ]

    # ─── DA (si module actif) ───
    da_data = {"count": 0, "amount": Decimal("0"), "closed": 0}
    try:
        from apps.da.models import PurchaseRequest
        if organization.has_module("da"):
            da_qs = PurchaseRequest.objects.filter(
                organization=organization,
                created_at__gte=start,
                created_at__lte=end,
            )
            da_valid = da_qs.filter(status__in=["VALIDEE", "COMMANDEE", "RECUE", "CLOTUREE", "RECUE_PARTIELLE"])
            da_data = {
                "count": da_qs.count(),
                "validated": da_valid.count(),
                "closed": da_qs.filter(status="CLOTUREE").count(),
                "amount": da_valid.aggregate(s=Sum("total_amount"))["s"] or Decimal("0"),
            }
    except Exception:
        pass

    # ─── Membres ───
    new_members = User.objects.filter(
        organization=organization,
        created_at__gte=start,
        created_at__lte=end,
    ).count()

    active_members = orders_qs.values("member_id").distinct().count()

    return {
        "year": year,
        "month": month,
        "period_label": f"{calendar.month_name[month]} {year}",
        "orders": orders_data,
        "top_products": top_products,
        "top_members": top_members,
        "da": da_data,
        "members": {
            "new": new_members,
            "active": active_members,
            "total": User.objects.filter(organization=organization).count(),
        },
    }


def generate_monthly_report_pdf(report):
    """Genere le PDF du rapport."""
    from xhtml2pdf import pisa

    data = collect_monthly_data(report.organization, report.year, report.month)

    html = render_to_string("reports/monthly_report_pdf.html", {
        "organization": report.organization,
        "year": report.year,
        "month": report.month,
        "data": data,
        "generated_at": timezone.now(),
    })

    buf = BytesIO()
    result = pisa.CreatePDF(html, dest=buf, encoding="utf-8")

    if result.err:
        raise Exception(f"xhtml2pdf error: {result.err}")

    filename = f"rapport-{report.organization.slug}-{report.year}-{report.month:02d}.pdf"
    report.pdf_file.save(filename, ContentFile(buf.getvalue()), save=False)

    return data


def generate_monthly_report_excel(report):
    """Genere l'Excel du rapport."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    data = collect_monthly_data(report.organization, report.year, report.month)

    wb = Workbook()

    # Feuille KPIs
    ws1 = wb.active
    ws1.title = "Synthese"
    ws1["A1"] = f"Rapport {data['period_label']} — {report.organization.name}"
    ws1["A1"].font = Font(bold=True, size=14, color="1E3A8A")

    ws1.append([])
    ws1.append(["Indicateur", "Valeur"])
    for cell in ws1[3]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2563EB")

    kpis = [
        ("Commandes totales", data["orders"]["total"]),
        ("Commandes validees", data["orders"]["validated"]),
        ("Commandes livrees", data["orders"]["delivered"]),
        ("Montant total (XOF)", float(data["orders"]["amount"])),
        ("Panier moyen (XOF)", float(data["orders"]["avg_amount"])),
        ("Nouveaux membres", data["members"]["new"]),
        ("Membres actifs", data["members"]["active"]),
    ]
    if data["da"]["count"]:
        kpis.append(("Demandes d'achat", data["da"]["count"]))
        kpis.append(("Montant DA (XOF)", float(data["da"]["amount"])))

    for label, val in kpis:
        ws1.append([label, val])

    ws1.column_dimensions["A"].width = 30
    ws1.column_dimensions["B"].width = 20

    # Feuille Top produits
    ws2 = wb.create_sheet("Top produits")
    ws2.append(["Produit", "Quantite", "CA (XOF)"])
    for cell in ws2[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2563EB")
    for p in data["top_products"]:
        ws2.append([p["name"], p["qty"], p["revenue"]])
    ws2.column_dimensions["A"].width = 40
    ws2.column_dimensions["B"].width = 15
    ws2.column_dimensions["C"].width = 20

    # Feuille Top membres
    ws3 = wb.create_sheet("Top membres")
    ws3.append(["Email", "Matricule", "Commandes", "Total (XOF)"])
    for cell in ws3[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2563EB")
    for m in data["top_members"]:
        ws3.append([m["email"], m["internal_id"], m["count"], m["total"]])
    ws3.column_dimensions["A"].width = 35
    ws3.column_dimensions["B"].width = 15
    ws3.column_dimensions["C"].width = 15
    ws3.column_dimensions["D"].width = 20

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"rapport-{report.organization.slug}-{report.year}-{report.month:02d}.xlsx"
    report.excel_file.save(filename, ContentFile(buf.read()), save=False)

    return data


def generate_report(year=None, month=None, organization=None):
    """Genere (ou regenere) un rapport mensuel."""
    from .models import MonthlyReport
    from apps.organizations.models import Organization

    now = timezone.now()
    year = year or now.year
    month = month or (now.month - 1 if now.month > 1 else 12)

    if organization is None:
        orgs = Organization.objects.filter(status__in=["ACTIVE", "TRIAL"])
    else:
        orgs = [organization]

    results = []
    for org in orgs:
        report, created = MonthlyReport.objects.get_or_create(
            organization=org, year=year, month=month,
            defaults={"status": "GENERATING"},
        )

        try:
            report.status = "GENERATING"
            report.save(update_fields=["status"])

            data = generate_monthly_report_pdf(report)
            generate_monthly_report_excel(report)

            report.orders_count = data["orders"]["total"]
            report.orders_amount = data["orders"]["amount"]
            report.da_count = data["da"]["count"]
            report.da_amount = data["da"]["amount"]
            report.members_active = data["members"]["active"]
            report.status = "READY"
            report.generated_at = timezone.now()
            report.error_message = ""
            report.save()

            results.append({"org": org.name, "status": "OK", "created": created})

        except Exception as e:
            report.status = "FAILED"
            report.error_message = str(e)[:500]
            report.save()
            results.append({"org": org.name, "status": "FAILED", "error": str(e)[:100]})

    return results


def send_report_email(report):
    """Envoie le rapport par email aux admins/DG."""
    from apps.accounts.models import User
    from apps.notifications.email_service import send_email

    if report.status != "READY" or not report.pdf_file:
        return 0

    recipients = User.objects.filter(
        organization=report.organization,
        role__in=["ADMIN", "GESTIONNAIRE"],
        is_active=True,
    ).filter(Q(is_dg=True) | Q(role="ADMIN"))

    sent = 0
    sent_to = []

    for user in recipients:
        if not user.email:
            continue

        success = send_email(
            recipient_email=user.email,
            subject=f"📊 Rapport mensuel {report.period_label}",
            template_name="emails/monthly_report.html",
            context={
                "organization": report.organization,
                "report": report,
                "user": user,
                "data": {
                    "orders_count": report.orders_count,
                    "orders_amount": report.orders_amount,
                    "da_count": report.da_count,
                },
            },
        )

        if success:
            sent += 1
            sent_to.append(user.email)

    if sent:
        report.sent_at = timezone.now()
        report.sent_to = sent_to
        report.save(update_fields=["sent_at", "sent_to"])

    return sent
